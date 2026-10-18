# services/document_export/xlsx_generator.py
"""
Excel (.xlsx) Generator for HEXA Enterprise Document Export System.
Uses openpyxl to generate structured, formatted Excel workbooks from DocumentModel.
"""

import os
from typing import List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services.document_export.models import DocumentModel, ExcelExportOptions


class XLSXGenerator:
    """Generates Excel workbooks from DocumentModel."""

    TEAL_FILL = PatternFill(start_color="0A9396", end_color="0A9396", fill_type="solid")
    NAVY_FILL = PatternFill(start_color="005F73", end_color="005F73", fill_type="solid")
    ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    SUMMARY_BG = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")

    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="005F73")
    SUBTITLE_FONT = Font(name="Calibri", size=11, bold=True, color="0A9396")
    BODY_FONT = Font(name="Calibri", size=11, color="2B2D42")
    META_FONT = Font(name="Calibri", size=9, italic=True, color="6C757D")

    THIN_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    @classmethod
    def generate(cls, doc_model: DocumentModel, options: ExcelExportOptions, output_path: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default blank sheet

        # 1. Summary Worksheet
        cls._create_summary_sheet(wb, doc_model, options)

        # 2. Data Table Worksheets
        all_tables = []
        for sec in doc_model.sections:
            all_tables.extend(sec.tables)
        for tbl in doc_model.tables:
            if tbl not in all_tables:
                all_tables.append(tbl)

        if all_tables:
            for idx, tbl in enumerate(all_tables, 1):
                sheet_title = f"Table {idx}"
                if tbl.caption:
                    safe_cap = str(tbl.caption)[:25].strip()
                    sheet_title = f"Data - {safe_cap}"
                cls._create_table_sheet(wb, sheet_title, tbl, options)
        else:
            # Create a general Data sheet from paragraphs if no explicit Markdown tables
            cls._create_text_data_sheet(wb, doc_model, options)

        # 3. References Worksheet
        if options.include_source_info and doc_model.references:
            cls._create_references_sheet(wb, doc_model.references, options)

        wb.save(output_path)
        return output_path

    @classmethod
    def _create_summary_sheet(cls, wb: openpyxl.Workbook, doc_model: DocumentModel, options: ExcelExportOptions):
        ws = wb.create_sheet(title="Summary")
        ws.views.sheetView[0].showGridLines = True

        # Header Block
        ws.cell(row=1, column=1, value="HEXA AI SOVEREIGN INDUSTRIAL WORKBENCH").font = cls.SUBTITLE_FONT
        ws.cell(row=2, column=1, value=doc_model.title).font = cls.TITLE_FONT

        meta = f"Author: {doc_model.author}  |  Date: {doc_model.created_at}  |  Source: {doc_model.source_type}"
        ws.cell(row=3, column=1, value=meta).font = cls.META_FONT

        row_cursor = 5

        # Executive Summary
        if doc_model.summary:
            ws.cell(row=row_cursor, column=1, value="Executive Summary").font = Font(name="Calibri", size=13, bold=True, color="0A9396")
            row_cursor += 1
            sum_cell = ws.cell(row=row_cursor, column=1, value=doc_model.summary)
            sum_cell.font = Font(name="Calibri", size=10.5, italic=True)
            sum_cell.fill = cls.SUMMARY_BG
            ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor+1, end_column=6)
            row_cursor += 3

        # Sections Overview
        ws.cell(row=row_cursor, column=1, value="Document Sections Overview").font = Font(name="Calibri", size=13, bold=True, color="005F73")
        row_cursor += 1

        headers = ["Section #", "Section Title", "Paragraphs", "Tables", "Code Blocks"]
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_cursor, column=c_idx, value=h)
            cell.fill = cls.TEAL_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        row_cursor += 1

        for idx, sec in enumerate(doc_model.sections, 1):
            ws.cell(row=row_cursor, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=row_cursor, column=2, value=sec.title or f"Section {idx}")
            ws.cell(row=row_cursor, column=3, value=len(sec.paragraphs)).alignment = Alignment(horizontal="center")
            ws.cell(row=row_cursor, column=4, value=len(sec.tables)).alignment = Alignment(horizontal="center")
            ws.cell(row=row_cursor, column=5, value=len(sec.code_blocks)).alignment = Alignment(horizontal="center")

            for col in range(1, 6):
                c = ws.cell(row=row_cursor, column=col)
                c.font = cls.BODY_FONT
                c.border = cls.THIN_BORDER
                if idx % 2 == 0:
                    c.fill = cls.ALT_FILL
            row_cursor += 1

        cls._auto_fit_columns(ws)

    @classmethod
    def _create_table_sheet(cls, wb: openpyxl.Workbook, sheet_title: str, tbl_model, options: ExcelExportOptions):
        # Truncate title to 31 chars (Excel limit)
        safe_title = sheet_title[:31].replace(":", "").replace("/", "").replace("\\", "")
        ws = wb.create_sheet(title=safe_title)
        ws.views.sheetView[0].showGridLines = True

        all_rows = []
        if tbl_model.headers:
            all_rows.append(tbl_model.headers)
        all_rows.extend(tbl_model.rows)

        if not all_rows:
            return

        for r_idx, row_data in enumerate(all_rows, start=1):
            is_header = (r_idx == 1 and bool(tbl_model.headers))
            for c_idx, val in enumerate(row_data, start=1):
                # Infer number/float if possible
                typed_val = cls._infer_type(val)
                cell = ws.cell(row=r_idx, column=c_idx, value=typed_val)
                cell.border = cls.THIN_BORDER

                if is_header:
                    cell.fill = cls.TEAL_FILL
                    cell.font = cls.HEADER_FONT
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = cls.BODY_FONT
                    if r_idx % 2 == 0:
                        cell.fill = cls.ALT_FILL

        if options.freeze_header and bool(tbl_model.headers):
            ws.freeze_panes = "A2"

        if options.auto_format_tables and ws.dimensions:
            ws.auto_filter.ref = ws.dimensions

        if options.auto_size_columns:
            cls._auto_fit_columns(ws)

    @classmethod
    def _create_text_data_sheet(cls, wb: openpyxl.Workbook, doc_model: DocumentModel, options: ExcelExportOptions):
        ws = wb.create_sheet(title="Extracted Data")
        ws.views.sheetView[0].showGridLines = True

        headers = ["Section", "Content Type", "Text Content"]
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c_idx, value=h)
            cell.fill = cls.TEAL_FILL
            cell.font = cls.HEADER_FONT

        r_cursor = 2
        for sec in doc_model.sections:
            sec_name = sec.title or "General"
            for p in sec.paragraphs:
                ws.cell(row=r_cursor, column=1, value=sec_name).font = cls.BODY_FONT
                ws.cell(row=r_cursor, column=2, value="Paragraph").font = cls.BODY_FONT
                ws.cell(row=r_cursor, column=3, value=p).font = cls.BODY_FONT
                r_cursor += 1

            for b_list in sec.bullet_lists:
                for item in b_list:
                    ws.cell(row=r_cursor, column=1, value=sec_name).font = cls.BODY_FONT
                    ws.cell(row=r_cursor, column=2, value="Bullet Point").font = cls.BODY_FONT
                    ws.cell(row=r_cursor, column=3, value=item).font = cls.BODY_FONT
                    r_cursor += 1

        if options.freeze_header:
            ws.freeze_panes = "A2"
        cls._auto_fit_columns(ws)

    @classmethod
    def _create_references_sheet(cls, wb: openpyxl.Workbook, references, options: ExcelExportOptions):
        ws = wb.create_sheet(title="References")
        ws.views.sheetView[0].showGridLines = True

        headers = ["#", "Source Name", "Page", "URL / Link", "Citation Notes"]
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c_idx, value=h)
            cell.fill = cls.TEAL_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        for idx, ref in enumerate(references, 1):
            row_idx = idx + 1
            ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=ref.source_name)
            ws.cell(row=row_idx, column=3, value=ref.page or "")
            ws.cell(row=row_idx, column=4, value=ref.url or "")
            ws.cell(row=row_idx, column=5, value=ref.citation or "")

            for col in range(1, 6):
                c = ws.cell(row=row_idx, column=col)
                c.font = cls.BODY_FONT
                c.border = cls.THIN_BORDER
                if idx % 2 == 0:
                    c.fill = cls.ALT_FILL

        if options.freeze_header:
            ws.freeze_panes = "A2"
        cls._auto_fit_columns(ws)

    @classmethod
    def _auto_fit_columns(cls, ws: openpyxl.worksheet.worksheet.Worksheet):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)

    @classmethod
    def _infer_type(cls, val: str):
        if not isinstance(val, str):
            return val
        v = val.strip()
        if not v:
            return ""
        # Try int
        try:
            return int(v)
        except ValueError:
            pass
        # Try float
        try:
            return float(v)
        except ValueError:
            pass
        return v
